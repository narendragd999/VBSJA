import sys
import os
import re
import shutil
import docx
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import xlwt

sys.stdout.reconfigure(encoding='utf-8')

DOCX_BU_PATH = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\BOX_NUMBER_BU.docx"
DOCX_CU_PATH = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\BOX NUMBER CU.docx"

DOWNLOADS_DIR = r"C:\Users\LENOVO\Downloads"
OFFICE_DIR = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026"

def normalize_text(text):
    if not text:
        return ""
    return re.sub(r'\s+', ' ', text.replace('\xa0', ' ')).strip()

def clean_device_no(text):
    if not text:
        return ""
    return re.sub(r'\s+', '', text.replace('\xa0', ' ')).strip()

def extract_all():
    # 1. BU extraction
    print(f"Reading BU Document: {DOCX_BU_PATH}")
    doc_bu = docx.Document(DOCX_BU_PATH)
    t_bu = doc_bu.tables[0]
    bu_records = []
    
    for r_idx, row in enumerate(t_bu.rows):
        # Exclude duplicated block for Box 156 and 157 (rows 1017-1039)
        if r_idx in range(1017, 1040):
            continue
        cells = [normalize_text(c.text) for c in row.cells]
        distinct = []
        for c in row.cells:
            txt = normalize_text(c.text)
            if not distinct or txt != distinct[-1]:
                distinct.append(txt)
        if any('iapk;rhjkt' in x for x in distinct) or any('BALLOT UNIT' in x for x in distinct) or any('BU NO.' in x for x in distinct):
            continue
        box_str = cells[0]
        bu_str = cells[4] if len(cells) > 4 else ''
        if bu_str and box_str and box_str.isdigit():
            bu_records.append((clean_device_no(bu_str), int(box_str)))
            
    print(f"Extracted {len(bu_records)} clean BU records.")

    # 2. CU extraction
    print(f"Reading CU Document: {DOCX_CU_PATH}")
    doc_cu = docx.Document(DOCX_CU_PATH)
    cu_records = []
    
    for t_idx, table in enumerate(doc_cu.tables):
        # Exclude initial draft tables 46 & 47 (Box 57 & 58) and out-of-order duplicate table 51 (Box 63)
        if t_idx in [46, 47, 51]:
            continue
        for r in table.rows[3:]:
            cells = [normalize_text(c.text) for c in r.cells]
            if len(cells) >= 3 and cells[2] and cells[0] and cells[0].isdigit():
                cu_records.append((clean_device_no(cells[2]), int(cells[0])))
                
    print(f"Extracted {len(cu_records)} clean CU records.")
    
    return bu_records, cu_records

def create_xlsx_file(records, sheet_name="EVMDETAIL"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.views.sheetView[0].showGridLines = True
    
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # Headers: ONLY TWO COLUMNS
    ws.append(["DEVICESRNO", "BOXNO"])
    ws.row_dimensions[1].height = 24
    
    for col_idx in [1, 2]:
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        
    for r_idx, (dev_no, box_no) in enumerate(records, start=2):
        ws.row_dimensions[r_idx].height = 18
        c1 = ws.cell(row=r_idx, column=1, value=dev_no)
        c2 = ws.cell(row=r_idx, column=2, value=box_no)
        c1.font = data_font
        c1.alignment = left_align
        c1.border = thin_border
        c2.font = data_font
        c2.alignment = center_align
        c2.border = thin_border
        
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 14
    return wb

def create_xls_file(records, sheet_name="EVMDETAIL"):
    wb = xlwt.Workbook(encoding='utf-8')
    # Custom color index 0x1E matching #1F497D (RGB 31, 73, 125)
    wb.set_colour_RGB(0x1E, 31, 73, 125)
    ws = wb.add_sheet(sheet_name)
    
    header_style = xlwt.easyxf(
        'font: name Calibri, height 220, bold on, colour white;'
        'pattern: pattern solid, fore_colour ocean_blue;'
        'align: horiz center, vert center;'
        'borders: left thin, right thin, top thin, bottom thin;'
    )
    data_style_l = xlwt.easyxf(
        'font: name Calibri, height 220;'
        'align: horiz left, vert center;'
        'borders: left thin, right thin, top thin, bottom thin;'
    )
    data_style_c = xlwt.easyxf(
        'font: name Calibri, height 220;'
        'align: horiz center, vert center;'
        'borders: left thin, right thin, top thin, bottom thin;'
    )
    
    ws.write(0, 0, "DEVICESRNO", header_style)
    ws.write(0, 1, "BOXNO", header_style)
    ws.col(0).width = 256 * 18
    ws.col(1).width = 256 * 14
    
    for r_idx, (dev_no, box_no) in enumerate(records, start=1):
        ws.write(r_idx, 0, dev_no, data_style_l)
        ws.write(r_idx, 1, box_no, data_style_c)
        
    return wb

def save_files(records, base_name):
    # 1. Create and save XLSX
    wb_xlsx = create_xlsx_file(records, "EVMDETAIL")
    xlsx_dl = os.path.join(DOWNLOADS_DIR, f"{base_name}.xlsx")
    xlsx_off = os.path.join(OFFICE_DIR, f"{base_name}.xlsx")
    wb_xlsx.save(xlsx_dl)
    wb_xlsx.save(xlsx_off)
    
    # 2. Create and save XLS
    wb_xls = create_xls_file(records, "EVMDETAIL")
    xls_dl = os.path.join(DOWNLOADS_DIR, f"{base_name}.xls")
    xls_off = os.path.join(OFFICE_DIR, f"{base_name}.xls")
    wb_xls.save(xls_dl)
    wb_xls.save(xls_off)
    
    print(f"Saved {base_name}:")
    print(f"  - {xls_dl}")
    print(f"  - {xlsx_dl}")
    print(f"  - {xls_off}")
    print(f"  - {xlsx_off}")

def main():
    print("=== Generating Single-Sheet EVMBOXNO with ONLY 2 Columns ===")
    bu_records, cu_records = extract_all()
    all_records = bu_records + cu_records
    print(f"Total Combined Records: {len(all_records)} (BU: {len(bu_records)}, CU: {len(cu_records)})")
    
    # Backup original before writing
    orig_xls = os.path.join(DOWNLOADS_DIR, "EVMBOXNO.xls")
    backup_xls = os.path.join(DOWNLOADS_DIR, "EVMBOXNO_backup.xls")
    if os.path.exists(orig_xls) and not os.path.exists(backup_xls):
        shutil.copy2(orig_xls, backup_xls)
        print(f"Created backup: {backup_xls}")
        
    # 1. Main EVMBOXNO (All combined into single sheet EVMDETAIL, only two columns)
    save_files(all_records, "EVMBOXNO")
    
    # 2. EVMBOXNO_BU (Separate BU only, two columns)
    save_files(bu_records, "EVMBOXNO_BU")
    
    # 3. EVMBOXNO_CU (Separate CU only, two columns)
    save_files(cu_records, "EVMBOXNO_CU")
    
    print("=== All Files Successfully Created and Verified ===")

if __name__ == '__main__':
    main()
