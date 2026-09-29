import sys
import os
import re
import shutil
import docx
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import win32com.client

sys.stdout.reconfigure(encoding='utf-8')

DOCX_BU_PATH = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\BOX_NUMBER_BU.docx"
DOCX_CU_PATH = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\BOX NUMBER CU.docx"

DOWNLOADS_DIR = r"C:\Users\LENOVO\Downloads"
OFFICE_DIR = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026"

def normalize_text(text):
    if not text:
        return ""
    return re.sub(r'\s+', ' ', text.replace('\xa0', ' ')).strip()

def clean_device_no(text):
    if not text:
        return ""
    return re.sub(r'\s+', '', text.replace('\xa0', ' ')).strip()

def extract_all():
    # 1. BU extraction
    print(f"Reading BU Document: {DOCX_BU_PATH}")
    doc_bu = docx.Document(DOCX_BU_PATH)
    t_bu = doc_bu.tables[0]
    bu_records = []
    
    for r_idx, row in enumerate(t_bu.rows):
        # Exclude duplicated block for Box 156 and 157
        if r_idx in range(1017, 1040):
            continue
        cells = [normalize_text(c.text) for c in row.cells]
        distinct = []
        for c in row.cells:
            txt = normalize_text(c.text)
            if not distinct or txt != distinct[-1]:
                distinct.append(txt)
        if any('iapk;rhjkt' in x for x in distinct) or any('BALLOT UNIT' in x for x in distinct) or any('BU NO.' in x for x in distinct):
            continue
        box_str = cells[0]
        bu_str = cells[4] if len(cells) > 4 else ''
        if bu_str and box_str and box_str.isdigit():
            bu_records.append((clean_device_no(bu_str), int(box_str)))
            
    print(f"Extracted {len(bu_records)} clean BU records.")

    # 2. CU extraction
    print(f"Reading CU Document: {DOCX_CU_PATH}")
    doc_cu = docx.Document(DOCX_CU_PATH)
    cu_records = []
    
    for t_idx, table in enumerate(doc_cu.tables):
        # Exclude initial draft tables 46 & 47 (Box 57 & 58) and out-of-order duplicate table 51 (Box 63)
        if t_idx in [46, 47, 51]:
            continue
        for r in table.rows[3:]:
            cells = [normalize_text(c.text) for c in r.cells]
            if len(cells) >= 3 and cells[2] and cells[0] and cells[0].isdigit():
                cu_records.append((clean_device_no(cells[2]), int(cells[0])))
                
    print(f"Extracted {len(cu_records)} clean CU records.")
    
    return bu_records, cu_records

def build_single_sheet_wb(records, sheet_name="EVMDETAIL"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.views.sheetView[0].showGridLines = True
    
    # Header styling matching uploaded image exactly (#1F497D dark blue, white bold text)
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # Headers: ONLY TWO COLUMNS
    ws.append(["DEVICESRNO", "BOXNO"])
    ws.row_dimensions[1].height = 24
    
    for col_idx in [1, 2]:
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        
    for r_idx, (dev_no, box_no) in enumerate(records, start=2):
        ws.row_dimensions[r_idx].height = 18
        c1 = ws.cell(row=r_idx, column=1, value=dev_no)
        c2 = ws.cell(row=r_idx, column=2, value=box_no)
        c1.font = data_font
        c1.alignment = left_align
        c1.border = thin_border
        c2.font = data_font
        c2.alignment = center_align
        c2.border = thin_border
        
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 14
    
    return wb

def save_both_formats(wb, base_name, excel_app):
    temp_xlsx = os.path.abspath(f"temp_{base_name}.xlsx")
    wb.save(temp_xlsx)
    
    # 1. Downloads
    xlsx_dl = os.path.join(DOWNLOADS_DIR, f"{base_name}.xlsx")
    xls_dl = os.path.join(DOWNLOADS_DIR, f"{base_name}.xls")
    shutil.copy2(temp_xlsx, xlsx_dl)
    
    doc = excel_app.Workbooks.Open(temp_xlsx)
    doc.SaveAs(os.path.abspath(xls_dl), 56) # 56 = xlExcel8 (.xls)
    doc.Close(SaveChanges=False)
    
    # 2. Office Dir
    os.makedirs(OFFICE_DIR, exist_ok=True)
    xlsx_off = os.path.join(OFFICE_DIR, f"{base_name}.xlsx")
    xls_off = os.path.join(OFFICE_DIR, f"{base_name}.xls")
    shutil.copy2(temp_xlsx, xlsx_off)
    
    doc = excel_app.Workbooks.Open(temp_xlsx)
    doc.SaveAs(os.path.abspath(xls_off), 56)
    doc.Close(SaveChanges=False)
    
    if os.path.exists(temp_xlsx):
        os.remove(temp_xlsx)
        
    print(f"Successfully generated {base_name}.xls and {base_name}.xlsx in Downloads and Office Data folders.")

def main():
    bu_records, cu_records = extract_all()
    all_records = bu_records + cu_records
    print(f"Total Combined Records (BU + CU): {len(all_records)}")
    
    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    
    try:
        # 1. Main File: EVMBOXNO (Combined BU + CU, 1,675 rows, single sheet EVMDETAIL, exactly two columns)
        print("\nCreating main EVMBOXNO (Combined BU + CU)...")
        wb_combined = build_single_sheet_wb(all_records, "EVMDETAIL")
        save_both_formats(wb_combined, "EVMBOXNO", excel)
        
        # 2. Separate BU File: EVMBOXNO_BU (Just in case BU upload is separate)
        print("\nCreating EVMBOXNO_BU...")
        wb_bu = build_single_sheet_wb(bu_records, "EVMDETAIL")
        save_both_formats(wb_bu, "EVMBOXNO_BU", excel)
        
        # 3. Separate CU File: EVMBOXNO_CU (Just in case CU upload is separate)
        print("\nCreating EVMBOXNO_CU...")
        wb_cu = build_single_sheet_wb(cu_records, "EVMDETAIL")
        save_both_formats(wb_cu, "EVMBOXNO_CU", excel)
        
    finally:
        excel.Quit()
        
    print("\nAll files generated successfully!")

if __name__ == '__main__':
    main()
