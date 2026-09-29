import sys
import os
import re
import shutil
import docx
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import win32com.client

sys.stdout.reconfigure(encoding='utf-8')

# Paths
DOCX_BU_PATH = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\BOX_NUMBER_BU.docx"
DOCX_CU_PATH = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\BOX NUMBER CU.docx"
TARGET_XLS_DOWNLOADS = r"C:\Users\LENOVO\Downloads\EVMBOXNO.xls"
BACKUP_XLS_DOWNLOADS = r"C:\Users\LENOVO\Downloads\EVMBOXNO_backup.xls"
TARGET_XLSX_DOWNLOADS = r"C:\Users\LENOVO\Downloads\EVMBOXNO.xlsx"

TARGET_XLS_OFFICE = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\EVMBOXNO.xls"
TARGET_XLSX_OFFICE = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\EVMBOXNO.xlsx"

def normalize_text(text):
    if not text:
        return ""
    return re.sub(r'\s+', ' ', text.replace('\xa0', ' ')).strip()

def clean_device_no(text):
    if not text:
        return ""
    # Remove all spaces as requested: e.g. 'MPB 207296' -> 'MPB207296'
    return re.sub(r'\s+', '', text.replace('\xa0', ' ')).strip()

def extract_bu_data():
    print(f"Reading BU Document: {DOCX_BU_PATH}")
    doc = docx.Document(DOCX_BU_PATH)
    t = doc.tables[0]
    
    raw_records = []
    clean_records = []
    duplicates_log = []
    
    for r_idx, row in enumerate(t.rows):
        cells = [normalize_text(c.text) for c in row.cells]
        distinct = []
        for c in row.cells:
            txt = normalize_text(c.text)
            if not distinct or txt != distinct[-1]:
                distinct.append(txt)
        
        # Skip header rows
        if any('iapk;rhjkt' in x for x in distinct) or any('BALLOT UNIT' in x for x in distinct) or any('BU NO.' in x for x in distinct):
            continue
            
        box_str = cells[0]
        sr_str = cells[2] if len(cells) > 2 else ''
        bu_str = cells[4] if len(cells) > 4 else ''
        
        if not bu_str or not box_str or not box_str.isdigit():
            continue
            
        box_no = int(box_str)
        sr_no = int(sr_str) if sr_str.isdigit() else sr_str
        dev_no = clean_device_no(bu_str)
        
        record = {
            'unit_type': 'BU',
            'devicesrno': dev_no,
            'boxno': box_no,
            'sr_no': sr_no,
            'orig_text': bu_str,
            'row_idx': r_idx
        }
        raw_records.append(record)
        
        # Rows 1017 to 1039 represent the duplicate paste of Box 156 and Box 157
        if r_idx in range(1017, 1040):
            duplicates_log.append({
                'unit_type': 'BU',
                'devicesrno': dev_no,
                'boxno': box_no,
                'sr_no': sr_no,
                'orig_text': bu_str,
                'location': f'Word Row {r_idx}',
                'status': 'Excluded from Clean BU',
                'remarks': f'Duplicate block paste of Box {box_no} (already recorded in Rows 991-1013)'
            })
        else:
            clean_records.append(record)
            
    print(f"BU: Extracted {len(raw_records)} raw records, {len(clean_records)} clean records, {len(duplicates_log)} duplicate records logged.")
    return raw_records, clean_records, duplicates_log

def extract_cu_data():
    print(f"Reading CU Document: {DOCX_CU_PATH}")
    doc = docx.Document(DOCX_CU_PATH)
    
    raw_records = []
    clean_records = []
    duplicates_log = []
    
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows[3:]):
            cells = [normalize_text(c.text) for c in row.cells]
            if len(cells) >= 3 and cells[2] and cells[0] and cells[0].isdigit():
                box_no = int(cells[0])
                sr_str = cells[1]
                sr_no = int(sr_str) if sr_str.isdigit() else sr_str
                cu_str = cells[2]
                dev_no = clean_device_no(cu_str)
                
                record = {
                    'unit_type': 'CU',
                    'devicesrno': dev_no,
                    'boxno': box_no,
                    'sr_no': sr_no,
                    'orig_text': cu_str,
                    'table_idx': t_idx,
                    'row_idx': r_idx + 3
                }
                raw_records.append(record)
                
                # Handling duplicates in CU:
                # Tables 46 & 47: earlier draft/incomplete entries for Box 57 & 58
                if t_idx in [46, 47]:
                    duplicates_log.append({
                        'unit_type': 'CU',
                        'devicesrno': dev_no,
                        'boxno': box_no,
                        'sr_no': sr_no,
                        'orig_text': cu_str,
                        'location': f'Word Table {t_idx} (Row {r_idx + 3})',
                        'status': 'Excluded from Clean CU',
                        'remarks': f'Initial draft/incomplete table for Box {box_no} (Replaced by corrected Tables 48 & 49)'
                    })
                # Table 51: out-of-order duplicate for Box 63 placed before Box 60
                elif t_idx == 51:
                    duplicates_log.append({
                        'unit_type': 'CU',
                        'devicesrno': dev_no,
                        'boxno': box_no,
                        'sr_no': sr_no,
                        'orig_text': cu_str,
                        'location': f'Word Table 51 (Row {r_idx + 3})',
                        'status': 'Excluded from Clean CU',
                        'remarks': 'Out-of-order duplicate table for Box 63 placed before Box 60 (retained at Table 55 in sequence)'
                    })
                else:
                    clean_records.append(record)
                    
    print(f"CU: Extracted {len(raw_records)} raw records, {len(clean_records)} clean records, {len(duplicates_log)} duplicate records logged.")
    return raw_records, clean_records, duplicates_log

def create_excel_workbook(bu_clean, cu_clean, bu_raw, cu_raw, all_duplicates):
    wb = openpyxl.Workbook()
    
    # Styles
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Dark Navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # 1. Sheet BU (Clean)
    ws_bu = wb.active
    ws_bu.title = "BU"
    ws_bu.views.sheetView[0].showGridLines = True
    ws_bu.append(["DEVICESRNO", "BOXNO"])
    for r in bu_clean:
        ws_bu.append([r['devicesrno'], r['boxno']])
        
    # 2. Sheet CU (Clean)
    ws_cu = wb.create_sheet(title="CU")
    ws_cu.views.sheetView[0].showGridLines = True
    ws_cu.append(["DEVICESRNO", "BOXNO"])
    for r in cu_clean:
        ws_cu.append([r['devicesrno'], r['boxno']])
        
    # 3. Sheet BU_Raw
    ws_bu_raw = wb.create_sheet(title="BU_Raw")
    ws_bu_raw.views.sheetView[0].showGridLines = True
    ws_bu_raw.append(["DEVICESRNO", "BOXNO", "SR_NO", "ORIGINAL_TEXT", "DOC_ROW"])
    for r in bu_raw:
        ws_bu_raw.append([r['devicesrno'], r['boxno'], r['sr_no'], r['orig_text'], r['row_idx']])
        
    # 4. Sheet CU_Raw
    ws_cu_raw = wb.create_sheet(title="CU_Raw")
    ws_cu_raw.views.sheetView[0].showGridLines = True
    ws_cu_raw.append(["DEVICESRNO", "BOXNO", "SR_NO", "ORIGINAL_TEXT", "TABLE_NO", "DOC_ROW"])
    for r in cu_raw:
        ws_cu_raw.append([r['devicesrno'], r['boxno'], r['sr_no'], r['orig_text'], r['table_idx'], r['row_idx']])
        
    # 5. Sheet Duplicates_Log
    ws_dup = wb.create_sheet(title="Duplicates_Log")
    ws_dup.views.sheetView[0].showGridLines = True
    ws_dup.append(["UNIT_TYPE", "DEVICESRNO", "BOXNO", "SR_NO", "ORIGINAL_TEXT", "LOCATION_IN_DOC", "STATUS_IN_CLEAN_SHEET", "REMARKS"])
    for r in all_duplicates:
        ws_dup.append([
            r['unit_type'],
            r['devicesrno'],
            r['boxno'],
            r['sr_no'],
            r['orig_text'],
            r['location'],
            r['status'],
            r['remarks']
        ])
        
    # Format all sheets
    for sheet in [ws_bu, ws_cu, ws_bu_raw, ws_cu_raw, ws_dup]:
        # Header formatting
        for col_idx in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
        
        # Row heights
        sheet.row_dimensions[1].height = 24
        
        # Data rows formatting and auto-width
        for row_idx in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row_idx].height = 18
            for col_idx in range(1, sheet.max_column + 1):
                c = sheet.cell(row=row_idx, column=col_idx)
                c.font = data_font
                c.border = thin_border
                # align box numbers and sr numbers center, text left
                if isinstance(c.value, int):
                    c.alignment = center_align
                else:
                    c.alignment = left_align
                    
        # Adjust column widths
        for col in sheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    return wb

def main():
    print("=== Starting EVM Box Number Extraction and Update ===")
    
    # 1. Extract BU & CU
    bu_raw, bu_clean, bu_dup = extract_bu_data()
    cu_raw, cu_clean, cu_dup = extract_cu_data()
    all_duplicates = bu_dup + cu_dup
    
    # Backup original EVMBOXNO.xls if exists
    if os.path.exists(TARGET_XLS_DOWNLOADS):
        print(f"Creating backup of original: {BACKUP_XLS_DOWNLOADS}")
        shutil.copy2(TARGET_XLS_DOWNLOADS, BACKUP_XLS_DOWNLOADS)
        
    # Create workbook
    wb = create_excel_workbook(bu_clean, cu_clean, bu_raw, cu_raw, all_duplicates)
    
    # Save temporary XLSX
    temp_xlsx = os.path.abspath("temp_evmboxno.xlsx")
    wb.save(temp_xlsx)
    print(f"Saved temporary workbook: {temp_xlsx}")
    
    # Save to Downloads and Office folders
    shutil.copy2(temp_xlsx, TARGET_XLSX_DOWNLOADS)
    print(f"Saved XLSX to Downloads: {TARGET_XLSX_DOWNLOADS}")
    
    os.makedirs(os.path.dirname(TARGET_XLSX_OFFICE), exist_ok=True)
    shutil.copy2(temp_xlsx, TARGET_XLSX_OFFICE)
    print(f"Saved XLSX to Office Data: {TARGET_XLSX_OFFICE}")
    
    # Convert to native .xls (Excel 97-2003 format 56) using Excel COM
    print("Converting to native .xls format via Microsoft Excel COM...")
    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        # Save to Downloads .xls
        doc = excel.Workbooks.Open(temp_xlsx)
        doc.SaveAs(os.path.abspath(TARGET_XLS_DOWNLOADS), 56)
        doc.Close(SaveChanges=False)
        print(f"Successfully generated native .xls: {TARGET_XLS_DOWNLOADS}")
        
        # Save to Office Data .xls
        doc = excel.Workbooks.Open(temp_xlsx)
        doc.SaveAs(os.path.abspath(TARGET_XLS_OFFICE), 56)
        doc.Close(SaveChanges=False)
        print(f"Successfully generated native .xls: {TARGET_XLS_OFFICE}")
    finally:
        excel.Quit()
        if os.path.exists(temp_xlsx):
            os.remove(temp_xlsx)
            
    print("=== Extraction, Cleaning and Generation Completed Successfully ===")

if __name__ == '__main__':
    main()
