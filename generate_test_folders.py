import sys
import os
import re
import shutil
import docx
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import xlwt

sys.stdout.reconfigure(encoding='utf-8')

DOCX_BU_PATH = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\BOX_NUMBER_BU.docx"
DOCX_CU_PATH = r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\BOX NUMBER CU.docx"

TEST_DIRS = [
    r"D:\Narendra Dangi\Office Data\ELECTION\FLC-PANCHAYAT-2026\TEST",
    r"C:\Users\LENOVO\Downloads\TEST",
    r"D:\Narendra Dangi\SCRIPTS\VBSJA\TEST"
]

def normalize_text(text):
    if not text:
        return ""
    return re.sub(r'\s+', ' ', text.replace('\xa0', ' ')).strip()

def clean_device_no(text):
    if not text:
        return ""
    return re.sub(r'\s+', '', text.replace('\xa0', ' ')).strip()

def extract_data():
    # 1. BU extraction
    print(f"Reading BU Document: {DOCX_BU_PATH}")
    doc_bu = docx.Document(DOCX_BU_PATH)
    t_bu = doc_bu.tables[0]
    bu_records = []
    
    for r_idx, row in enumerate(t_bu.rows):
        if r_idx in range(1017, 1040): # Duplicate paste of Box 156 & 157
            continue
        cells = [normalize_text(c.text) for c in row.cells]
        distinct = []
        for c in row.cells:
            txt = normalize_text(c.text)
            if not distinct or txt != distinct[-1]:
                distinct.append(txt)
        if any('iapk;rhjkt' in x for x in distinct) or any('BALLOT UNIT' in x for x in distinct) or any('BU NO.' in x for x in distinct):
            continue
        box_str = cells[0]
        bu_str = cells[4] if len(cells) > 4 else ''
        if bu_str and box_str and box_str.isdigit():
            bu_records.append((clean_device_no(bu_str), int(box_str)))
            
    print(f"BU: Extracted {len(bu_records)} clean records.")

    # 2. CU extraction
    print(f"Reading CU Document: {DOCX_CU_PATH}")
    doc_cu = docx.Document(DOCX_CU_PATH)
    cu_records = []
    
    for t_idx, table in enumerate(doc_cu.tables):
        if t_idx in [46, 47, 51]: # Draft Box 57/58 and duplicate Box 63
            continue
        for r in table.rows[3:]:
            cells = [normalize_text(c.text) for c in r.cells]
            if len(cells) >= 3 and cells[2] and cells[0] and cells[0].isdigit():
                cu_records.append((clean_device_no(cells[2]), int(cells[0])))
                
    print(f"CU: Extracted {len(cu_records)} clean records.")
    return bu_records, cu_records

def create_single_sheet_xlsx(records, sheet_name):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.views.sheetView[0].showGridLines = True
    
    # Header style matching uploaded image: Dark Blue #1F497D, White Bold Text
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    ws.append(["DEVICESRNO", "BOXNO"])
    ws.row_dimensions[1].height = 24
    
    for col_idx in [1, 2]:
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        
    for r_idx, (dev_no, box_no) in enumerate(records, start=2):
        ws.row_dimensions[r_idx].height = 18
        c1 = ws.cell(row=r_idx, column=1, value=dev_no)
        c2 = ws.cell(row=r_idx, column=2, value=box_no)
        c1.font = data_font
        c1.alignment = left_align
        c1.border = thin_border
        c2.font = data_font
        c2.alignment = center_align
        c2.border = thin_border
        
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 14
    return wb

def create_single_sheet_xls(records, sheet_name):
    wb = xlwt.Workbook(encoding='utf-8')
    wb.set_colour_RGB(0x1E, 31, 73, 125) # Dark Blue #1F497D
    ws = wb.add_sheet(sheet_name)
    
    header_style = xlwt.easyxf(
        'font: name Calibri, height 220, bold on, colour white;'
        'pattern: pattern solid, fore_colour ocean_blue;'
        'align: horiz center, vert center;'
        'borders: left thin, right thin, top thin, bottom thin;'
    )
    data_style_l = xlwt.easyxf(
        'font: name Calibri, height 220;'
        'align: horiz left, vert center;'
        'borders: left thin, right thin, top thin, bottom thin;'
    )
    data_style_c = xlwt.easyxf(
        'font: name Calibri, height 220;'
        'align: horiz center, vert center;'
        'borders: left thin, right thin, top thin, bottom thin;'
    )
    
    ws.write(0, 0, "DEVICESRNO", header_style)
    ws.write(0, 1, "BOXNO", header_style)
    ws.col(0).width = 256 * 18
    ws.col(1).width = 256 * 14
    
    for r_idx, (dev_no, box_no) in enumerate(records, start=1):
        ws.write(r_idx, 0, dev_no, data_style_l)
        ws.write(r_idx, 1, box_no, data_style_c)
        
    return wb

def save_to_all_test_dirs(records, file_basename, sheet_name):
    # Build in-memory / temp workbooks
    wb_xlsx = create_single_sheet_xlsx(records, sheet_name)
    wb_xls = create_single_sheet_xls(records, sheet_name)
    
    for base_dir in TEST_DIRS:
        os.makedirs(base_dir, exist_ok=True)
        xlsx_path = os.path.join(base_dir, f"{file_basename}.xlsx")
        xls_path = os.path.join(base_dir, f"{file_basename}.xls")
        
        wb_xlsx.save(xlsx_path)
        wb_xls.save(xls_path)
        print(f"Saved: {xls_path} (1 sheet: '{sheet_name}', {len(records)} rows)")

def main():
    print("=== Generating Single-Sheet Separate BU and CU Files in TEST Folder ===")
    bu_records, cu_records = extract_data()
    
    # 1. BU Files (Single Sheet, exactly 2 columns: DEVICESRNO, BOXNO)
    # With sheet name 'EVMDETAIL' (portal standard)
    save_to_all_test_dirs(bu_records, "EVMBOXNO_BU", "EVMDETAIL")
    save_to_all_test_dirs(bu_records, "BU", "BU")
    
    # 2. CU Files (Single Sheet, exactly 2 columns: DEVICESRNO, BOXNO)
    # With sheet name 'EVMDETAIL' (portal standard)
    save_to_all_test_dirs(cu_records, "EVMBOXNO_CU", "EVMDETAIL")
    save_to_all_test_dirs(cu_records, "CU", "CU")
    
    # 3. Also create dedicated BU and CU subfolders where the file is simply named EVMBOXNO.xls / EVMBOXNO.xlsx
    for base_dir in TEST_DIRS:
        bu_folder = os.path.join(base_dir, "BU")
        cu_folder = os.path.join(base_dir, "CU")
        os.makedirs(bu_folder, exist_ok=True)
        os.makedirs(cu_folder, exist_ok=True)
        
        # In BU folder: EVMBOXNO.xls with sheet EVMDETAIL (only 1 sheet)
        wb_bu_xlsx = create_single_sheet_xlsx(bu_records, "EVMDETAIL")
        wb_bu_xls = create_single_sheet_xls(bu_records, "EVMDETAIL")
        wb_bu_xlsx.save(os.path.join(bu_folder, "EVMBOXNO.xlsx"))
        wb_bu_xls.save(os.path.join(bu_folder, "EVMBOXNO.xls"))
        
        # In CU folder: EVMBOXNO.xls with sheet EVMDETAIL (only 1 sheet)
        wb_cu_xlsx = create_single_sheet_xlsx(cu_records, "EVMDETAIL")
        wb_cu_xls = create_single_sheet_xls(cu_records, "EVMDETAIL")
        wb_cu_xlsx.save(os.path.join(cu_folder, "EVMBOXNO.xlsx"))
        wb_cu_xls.save(os.path.join(cu_folder, "EVMBOXNO.xls"))

    print("=== Successfully Created All Single-Sheet Files in TEST Folders ===")

if __name__ == '__main__':
    main()
